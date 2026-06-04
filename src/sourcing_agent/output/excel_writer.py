from __future__ import annotations

import os
from collections import defaultdict
from typing import TYPE_CHECKING, Any

from loguru import logger
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    from ..config import Config
    from ..models import PaperRecord

_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF")

_ALL_PAPERS_COLS = [
    "title",
    "authors",
    "year",
    "venue",
    "doi",
    "arxiv_id",
    "primary_section",
    "secondary_sections",
    "kunze_dimension",
    "technique_cluster",
    "cluster_assignment_confidence",
    "source_type",
    "section_fit_score",
    "contribution_score",
    "recency_score",
    "weighted_score",
    "inclusion_decision",
    "citation_count",
    "open_access_pdf_url",
    "source_database",
    "is_cross_cutting",
    "is_historical",
    "space_heritage",
    "deployment_constraints_discussed",
    "compute_requirements_noted",
    "radiation_robustness_discussed",
    "operational_description_present",
    "alignment_paper",
    "compositional_alignment",
    "general_alignment_robotics_context",
    "mission_or_system_name",
    "agent_notes",
    "abstract_snippet",
]

_KUNZE_DIMS = ["D1", "D2", "D3", "D4", "D5", "D6", "alignment"]
_S4_SECTIONS = [
    "S4a:navigation",
    "S4b:perception",
    "S4c:reasoning",
    "S4d:planning",
    "S4e:interaction",
    "S4f:learning",
    "S4g:alignment",
]


def write_to_excel(records: list[PaperRecord], config: Config) -> None:
    path = config.output.excel.path
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)

    wb = Workbook()

    _write_all_papers(wb, records, config)
    _write_section_coverage(wb, records, config)
    _write_analytical_data(wb, records, config)
    _write_prisma(wb, config)

    # Remove default empty sheet
    if "Sheet" in wb.sheetnames:
        del wb["Sheet"]

    wb.save(path)
    logger.info(f"Excel written: {path} ({len(records)} total records, 4 sheets)")


# ── Sheet 1 — All Papers ──────────────────────────────────────────────────────


def _write_all_papers(wb: Workbook, records: list[PaperRecord], config: Config) -> None:
    ws = wb.create_sheet("All Papers")
    _write_header(ws, _ALL_PAPERS_COLS)

    for row_idx, record in enumerate(records, start=2):
        fill = _row_fill(record.inclusion_decision)
        values = _record_to_row(record)
        for col_idx, val in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if fill:
                cell.fill = fill

    _auto_width(ws, _ALL_PAPERS_COLS)


def _record_to_row(record: PaperRecord) -> list:
    return [
        record.title,
        "; ".join(record.authors[:5]) + (" et al." if len(record.authors) > 5 else ""),
        record.year,
        record.venue,
        record.doi,
        record.arxiv_id,
        record.primary_section,
        "; ".join(record.secondary_sections),
        record.kunze_dimension,
        record.technique_cluster,
        record.cluster_assignment_confidence,
        record.source_type,
        record.section_fit_score,
        record.contribution_score,
        record.recency_score,
        record.weighted_score,
        record.inclusion_decision,
        record.citation_count,
        record.open_access_pdf_url,
        record.source_database,
        record.is_cross_cutting,
        record.is_historical,
        record.space_heritage,
        record.deployment_constraints_discussed,
        record.compute_requirements_noted,
        record.radiation_robustness_discussed,
        record.operational_description_present,
        record.alignment_paper,
        record.compositional_alignment,
        record.general_alignment_robotics_context,
        record.mission_or_system_name,
        record.agent_notes,
        record.abstract_snippet,
    ]


# ── Sheet 2 — Section Coverage ────────────────────────────────────────────────


def _write_section_coverage(
    wb: Workbook, records: list[PaperRecord], config: Config
) -> None:
    ws = wb.create_sheet("Section Coverage")
    headers = [
        "section_tag",
        "section_title",
        "kunze_dimension",
        "target_paper_count",
        "included_count",
        "maybe_count",
        "coverage_status",
        "cluster_confidence_high_pct",
    ]
    _write_header(ws, headers)

    included = [r for r in records if r.inclusion_decision == "include"]
    maybe = [r for r in records if r.inclusion_decision == "maybe"]

    for row_idx, (_tag, sec) in enumerate(config.paper_structure.items(), start=2):
        inc_count = sum(
            1
            for r in included
            if sec.section_tag in [r.primary_section, *r.secondary_sections]
        )
        may_count = sum(
            1
            for r in maybe
            if sec.section_tag in [r.primary_section, *r.secondary_sections]
        )
        target_min = sec.target_min
        coverage = "adequate" if inc_count >= target_min * 0.5 else "needs more"

        high_conf_pct = ""
        if sec.section_tag.startswith("S4"):
            s4_papers = [
                r
                for r in included
                if sec.section_tag in [r.primary_section, *r.secondary_sections]
            ]
            if s4_papers:
                high = sum(
                    1 for r in s4_papers if r.cluster_assignment_confidence == "high"
                )
                high_conf_pct = f"{round(high / len(s4_papers) * 100)}%"

        row = [
            sec.section_tag,
            sec.title,
            sec.kunze_dimension or "",
            sec.target_paper_count,
            inc_count,
            may_count,
            coverage,
            high_conf_pct,
        ]
        for col_idx, val in enumerate(row, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            if coverage == "needs more" and col_idx == 7:
                cell.fill = _YELLOW

    _auto_width(ws, headers)


# ── Sheet 3 — Analytical Data ─────────────────────────────────────────────────


def _write_analytical_data(
    wb: Workbook, records: list[PaperRecord], config: Config
) -> None:
    ws = wb.create_sheet("Analytical Data")
    row = 1

    # Tab 3a — S2 Mission Taxonomy
    ws.cell(row=row, column=1, value="TAB 3a — S2 Mission Taxonomy").font = Font(
        bold=True, size=12
    )
    row += 1
    mission_headers = [
        "mission_name",
        "year_deployed",
        "D1_nav",
        "D2_perc",
        "D3_kr",
        "D4_plan",
        "D5_interact",
        "D6_learn",
        "reference_paper",
    ]
    _write_header_at(ws, row, mission_headers)
    row += 1

    s2_included = [
        r
        for r in records
        if r.inclusion_decision in ("include", "maybe")
        and r.primary_section
        and "S2" in r.primary_section
        and r.mission_or_system_name
    ]
    missions_seen: dict[str, dict] = {}
    for r in s2_included:
        mn = r.mission_or_system_name or ""
        if mn not in missions_seen:
            missions_seen[mn] = {
                "year": r.year,
                "D1": "",
                "D2": "",
                "D3": "",
                "D4": "",
                "D5": "",
                "D6": "",
                "ref": r.title[:80],
            }
        if r.kunze_dimension:
            dim = r.kunze_dimension
            key = dim.replace("D", "D") if dim.startswith("D") else ""
            if key and key in missions_seen[mn]:
                missions_seen[mn][key] = "✓"

    for mn, data in missions_seen.items():
        mission_row = [
            mn,
            data["year"],
            data["D1"],
            data["D2"],
            data["D3"],
            data["D4"],
            data["D5"],
            data["D6"],
            data["ref"],
        ]
        for col_idx, val in enumerate(mission_row, start=1):
            ws.cell(row=row, column=col_idx, value=val)
        row += 1

    row += 2

    # Tab 3b — S2 Deployment Lag
    ws.cell(row=row, column=1, value="TAB 3b — S2 Deployment Lag").font = Font(
        bold=True, size=12
    )
    row += 1
    lag_headers = [
        "ai_capability",
        "terrestrial_first_pub_year",
        "space_deployment_year",
        "deployment_lag_years",
        "evidence_paper",
    ]
    _write_header_at(ws, row, lag_headers)
    row += 1

    for r in s2_included:
        if r.space_heritage and r.technique_cluster:
            lag_row = [
                r.technique_cluster,
                r.year or "",
                "",  # space_deployment_year — filled at writing stage
                "",  # lag — computed at writing stage
                r.title[:80],
            ]
            for col_idx, val in enumerate(lag_row, start=1):
                ws.cell(row=row, column=col_idx, value=val)
            row += 1

    row += 2

    # Tab 3c — S4 Technique Cluster Summary
    ws.cell(row=row, column=1, value="TAB 3c — S4 Technique Cluster Summary").font = (
        Font(bold=True, size=12)
    )
    row += 1
    cluster_headers = [
        "kunze_dimension",
        "technique_cluster",
        "paper_count",
        "space_heritage_count",
        "deployment_constraints_count",
        "compute_noted_count",
        "radiation_noted_count",
        "high_confidence_pct",
    ]
    _write_header_at(ws, row, cluster_headers)
    row += 1

    s4_included = [
        r
        for r in records
        if r.inclusion_decision in ("include", "maybe") and r.kunze_dimension
    ]
    cluster_stats: dict[str, dict] = defaultdict(
        lambda: {
            "dim": "",
            "count": 0,
            "space": 0,
            "deploy": 0,
            "compute": 0,
            "rad": 0,
            "high": 0,
        }
    )
    for r in s4_included:
        if r.technique_cluster:
            c = cluster_stats[r.technique_cluster]
            c["dim"] = r.kunze_dimension or ""
            c["count"] += 1
            if r.space_heritage:
                c["space"] += 1
            if r.deployment_constraints_discussed:
                c["deploy"] += 1
            if r.compute_requirements_noted:
                c["compute"] += 1
            if r.radiation_robustness_discussed:
                c["rad"] += 1
            if r.cluster_assignment_confidence == "high":
                c["high"] += 1

    for cluster, stats in sorted(cluster_stats.items(), key=lambda x: x[1]["dim"]):
        pct = (
            f"{round(stats['high'] / max(stats['count'], 1) * 100)}%"
            if stats["count"]
            else ""
        )
        cr = [
            stats["dim"],
            cluster,
            stats["count"],
            stats["space"],
            stats["deploy"],
            stats["compute"],
            stats["rad"],
            pct,
        ]
        for col_idx, val in enumerate(cr, start=1):
            ws.cell(row=row, column=col_idx, value=val)
        row += 1

    row += 2

    # Tab 3d — Bibliometric Trend
    ws.cell(
        row=row, column=1, value="TAB 3d — Bibliometric Trend (included S4 papers)"
    ).font = Font(bold=True, size=12)
    row += 1
    years = list(range(2015, 2027))
    trend_headers = ["year"] + [
        f"S4{x}_count"
        for x in [
            "a_nav",
            "b_perc",
            "c_reason",
            "d_plan",
            "e_interact",
            "f_learn",
            "g_align",
        ]
    ]
    _write_header_at(ws, row, trend_headers)
    row += 1

    section_map = {
        "S4a:navigation": 1,
        "S4b:perception": 2,
        "S4c:reasoning": 3,
        "S4d:planning": 4,
        "S4e:interaction": 5,
        "S4f:learning": 6,
        "S4g:alignment": 7,
    }
    trend: dict[int, list[int]] = {y: [0] * 7 for y in years}

    for r in s4_included:
        if r.year and r.year in trend and r.primary_section:
            idx = section_map.get(r.primary_section)
            if idx:
                trend[r.year][idx - 1] += 1

    for year in years:
        yr = [year] + trend[year]
        for col_idx, val in enumerate(yr, start=1):
            ws.cell(row=row, column=col_idx, value=val)
        row += 1

    _auto_width(ws, trend_headers)


# ── Sheet 4 — PRISMA Flow ─────────────────────────────────────────────────────


def _write_prisma(wb: Workbook, config: Config) -> None:
    ws = wb.create_sheet("PRISMA Flow")
    headers = ["stage", "database", "count", "notes"]
    _write_header(ws, headers)
    # Actual PRISMA data is populated from progress tracker at run time.
    # Write a placeholder row so the sheet isn't empty.
    ws.cell(row=2, column=1, value="[Populated from outputs/sourcing-progress.txt]")
    _auto_width(ws, headers)


# ── Helpers ───────────────────────────────────────────────────────────────────


def _write_header(ws: Any, headers: list[str], row: int = 1) -> None:
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center")


def _write_header_at(ws: Any, row: int, headers: list[str]) -> None:
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT


def _row_fill(decision: str | None) -> PatternFill | None:
    if decision == "include":
        return _GREEN
    if decision == "maybe":
        return _YELLOW
    if decision == "exclude":
        return _RED
    return None


def _auto_width(ws: Any, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = min(max(len(header) + 4, 12), 50)
